"""Create a generic research-landscape PPTX from the local catalog and snapshot."""

from __future__ import annotations

import json
from collections import Counter
from pathlib import Path

from common import ANALYSIS, ROOT, clean, load_papers, paper_d, publication, save_json


def require_pptx():
    try:
        from pptx import Presentation
        from pptx.dml.color import RGBColor
        from pptx.enum.text import PP_ALIGN
        from pptx.util import Inches, Pt
    except ImportError as exc:
        raise SystemExit("生成 PPTX 需要 python-pptx；请先运行: pip install -r requirements.txt") from exc
    return Presentation, RGBColor, PP_ALIGN, Inches, Pt


def add_text(slide, text, left, top, width, height, *, size=18, color=None, bold=False):
    _, RGBColor, _, Inches, Pt = require_pptx()
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.word_wrap = True
    frame.clear()
    paragraph = frame.paragraphs[0]
    paragraph.text = str(text)
    paragraph.font.size = Pt(size)
    paragraph.font.bold = bold
    if color:
        paragraph.font.color.rgb = RGBColor(*color)
    return box


def main() -> int:
    Presentation, RGBColor, PP_ALIGN, Inches, Pt = require_pptx()
    papers = load_papers()
    grouped_path = ANALYSIS / "research-landscape-grouped.json"
    grouped = json.loads(grouped_path.read_text(encoding="utf-8")) if grouped_path.exists() else {"dAggregation": [], "groups": [], "unassigned": []}
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    ink, muted, blue, teal = (22, 32, 42), (93, 107, 120), (23, 105, 170), (0, 139, 139)

    slide = prs.slides.add_slide(blank)
    add_text(slide, "Computer Research Landscape", 0.7, 0.7, 11.8, 0.7, size=30, color=blue, bold=True)
    add_text(slide, f"论文 {len(papers)} 篇 · 已分组 {grouped.get('assignedCount', 0)} · 未分组 {grouped.get('unassignedCount', 0)}", 0.75, 1.6, 11, 0.45, size=18, color=muted)
    add_text(slide, "研究对象 → Q → 论文 → 论文级 D → 跨 Q D 聚合", 0.75, 2.35, 11, 0.5, size=22, color=ink)
    add_text(slide, "本演示文稿由本项目本地数据生成。", 0.75, 6.55, 11, 0.3, size=12, color=muted)

    years = Counter(publication(paper.get("venueYear"))[1] for paper in papers)
    slide = prs.slides.add_slide(blank)
    add_text(slide, "年份分布", 0.7, 0.55, 11.8, 0.5, size=26, color=blue, bold=True)
    y = 1.35
    for year, count in sorted(years.items()):
        add_text(slide, year, 1.0, y, 1.4, 0.35, size=16, color=ink, bold=True)
        bar = slide.shapes.add_shape(1, Inches(2.0), Inches(y + 0.03), Inches(max(0.15, count * 0.35)), Inches(0.28))
        bar.fill.solid(); bar.fill.fore_color.rgb = RGBColor(*teal); bar.line.fill.background()
        add_text(slide, str(count), 2.1 + max(0.15, count * 0.35), y, 0.7, 0.35, size=14, color=muted)
        y += 0.55

    slide = prs.slides.add_slide(blank)
    add_text(slide, "论文级 D 的跨 Q 聚合", 0.7, 0.55, 11.8, 0.5, size=26, color=blue, bold=True)
    add_text(slide, "聚合来自论文记录中的 D，不从预设方向推断。", 0.75, 1.05, 11.5, 0.35, size=14, color=muted)
    entries = grouped.get("dAggregation", [])[:12]
    y = 1.55
    for entry in entries:
        label = clean(entry.get("label"))
        add_text(slide, label, 0.9, y, 5.6, 0.34, size=14, color=ink)
        count = len(entry.get("paperIds", []))
        bar = slide.shapes.add_shape(1, Inches(6.5), Inches(y + 0.04), Inches(max(0.15, min(5.2, count * 0.55))), Inches(0.25))
        bar.fill.solid(); bar.fill.fore_color.rgb = RGBColor(*teal); bar.line.fill.background()
        add_text(slide, f"{count} 篇", 11.9, y, 0.8, 0.3, size=13, color=muted)
        y += 0.43

    groups = grouped.get("groups", [])
    slide = prs.slides.add_slide(blank)
    add_text(slide, "Q 分组概览", 0.7, 0.55, 11.8, 0.5, size=26, color=blue, bold=True)
    y = 1.25
    for group in groups[:14]:
        count = len(group.get("papers", [])) + sum(len(child.get("papers", [])) for child in group.get("children", []))
        add_text(slide, f"{group.get('name', '未命名')}（{count}）", 0.95, y, 11.0, 0.35, size=17, color=ink, bold=True)
        y += 0.45
    if not groups:
        add_text(slide, "尚未建立人工 Q 分组；论文仍保留在未分组目录。", 0.95, y, 11, 0.4, size=16, color=muted)

    for offset in range(0, len(papers), 8):
        batch = papers[offset:offset + 8]
        slide = prs.slides.add_slide(blank)
        add_text(slide, f"论文记录 {offset + 1}–{offset + len(batch)}", 0.7, 0.55, 11.8, 0.5, size=24, color=blue, bold=True)
        y = 1.2
        for paper in batch:
            _, year, venue = publication(paper.get("venueYear"))
            d = paper_d(paper).get("primary", "未报告")
            title = clean(paper.get("title"))
            add_text(slide, f"{year} · {venue} · {title}", 0.85, y, 11.5, 0.28, size=12, color=ink, bold=True)
            add_text(slide, f"Q：{clean(paper.get('Q'))}｜D：{d}", 1.05, y + 0.27, 11.1, 0.34, size=11, color=muted)
            y += 0.68

    out = ROOT / "analysis" / "research-landscape.pptx"
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out)
    print(json.dumps({"pptx": str(out), "slides": len(prs.slides), "papers": len(papers)}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
